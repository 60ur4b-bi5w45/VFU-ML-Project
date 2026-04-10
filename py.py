from openpyxl import Workbook

wb = Workbook()

# Sheet 1: Pi 2D
ws1 = wb.active
ws1.title = "Pi_2D"

ws1["A1"] = "x"
ws1["B1"] = "y"
ws1["C1"] = "Inside Circle (1/0)"

for i in range(2, 10002):
    ws1[f"A{i}"] = "=RAND()"
    ws1[f"B{i}"] = "=RAND()"
    ws1[f"C{i}"] = f"=IF(A{i}^2+B{i}^2<=1,1,0)"

ws1["E1"] = "Estimated Pi"
ws1["E2"] = "=4*AVERAGE(C2:C10001)"

# Sheet 2: Pi 3D
ws2 = wb.create_sheet("Pi_3D")

ws2["A1"] = "x"
ws2["B1"] = "y"
ws2["C1"] = "z"
ws2["D1"] = "Inside Sphere (1/0)"

for i in range(2, 10002):
    ws2[f"A{i}"] = "=RAND()"
    ws2[f"B{i}"] = "=RAND()"
    ws2[f"C{i}"] = "=RAND()"
    ws2[f"D{i}"] = f"=IF(A{i}^2+B{i}^2+C{i}^2<=1,1,0)"

ws2["F1"] = "Estimated Pi"
ws2["F2"] = "=6*AVERAGE(D2:D10001)"

# Sheet 3: System 1
ws3 = wb.create_sheet("System1")

headers = ["X1(10)", "X2(8)", "X3(7)", "X4(5)", "X5(6)", "System Failure Time", "Survive >=7"]
for col, h in enumerate(headers, 1):
    ws3.cell(row=1, column=col, value=h)

for i in range(2, 10002):
    ws3[f"A{i}"] = "=-10*LN(RAND())"
    ws3[f"B{i}"] = "=-8*LN(RAND())"
    ws3[f"C{i}"] = "=-7*LN(RAND())"
    ws3[f"D{i}"] = "=-5*LN(RAND())"
    ws3[f"E{i}"] = "=-6*LN(RAND())"
    ws3[f"F{i}"] = f"=MIN(A{i}:E{i})"
    ws3[f"G{i}"] = f"=IF(F{i}>=7,1,0)"

ws3["I1"] = "Estimated p"
ws3["I2"] = "=AVERAGE(G2:G10001)"

ws3["I4"] = "Required n"
ws3["I5"] = "=(1.96^2*I2*(1-I2))/(0.01^2)"

# Save file
file_path = "/mnt/data/Monte_Carlo_Assignment.xlsx"
wb.save(file_path)

file_path